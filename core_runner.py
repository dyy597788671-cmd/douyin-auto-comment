#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PhoneWorkbench 1.0: strategy/state engine; no mobile SDK or touch code.

Requires Python >= 3.8 and openpyxl >= 3.1.
All processes sharing a project must use this gateway on the same host.
Excel is an input/output table, not a database; close it while running jobs.

Project files:
  config/settings.json              optional, defaults below
  data/task_cases.xlsx              required; seven headers in HEADERS,
                                   plus Expected_Douyin_ID for profile ID checks
                                   Content_Type (视频/图文) for search/detail layout,
                                   Comments: one candidate per newline in each project
  data/author_whitelist.json        optional: {"authors": [
      {"uid": "123", "nicknames": ["测试作者"]}]}
  data/runtime_state.json           managed; do not edit/delete while running
  logs/execution.log
  execution_data/执行数据_<账号>.csv   completed action counts for this account
  execution_data/project_progress.json  assigned devices, per-account results and totals
  workspaces/<account>/<claim>/         isolated screenshot/OCR files

Target_Device_Count is the requested participant count. For these projects,
allocate_task(task_id, connected_devices) must first receive ShadowBot's actual
connection snapshot; settings.sender_accounts is never a connection inventory.
Workers must pass their real physical_device_id when claiming/restoring tasks.
The ten-column legacy input remains readable without resetting prior progress.

Initialize an empty project: python core_runner.py --root PATH --init
Call through Python: gateway = UnifiedGateway(PATH, device_id="phone-01")
task = gateway.get_next_task(); plan = gateway.generate_strategy()
Pass task["run_token"] to record_result(..., run_token=...).
An active job survives a crash; get_active_task() retrieves its context,
but never automatically replays UI actions of uncertain outcome.
Numeric Scheduled_Time means minutes AFTER its first persisted scan.
Time-only schedules use the current system-local calendar date.
Naive timestamps use timezone_offset_minutes (default UTC+08:00).
Nickname MATCH means exact normalized nickname equality, not unique identity.
Profile Douyin IDs use {"douyin_id": ...}; these are separate from internal uid.
Missing/truncated Douyin IDs return UNCERTAIN without nickname fallback.
Legacy {"uid": ...} and nickname-only assertions retain their existing behavior.
Cooldown settings default to 300..600 seconds and allow zero for local debugging.
Exceptions expose .code; the CLI returns a JSON error envelope and exit code 1.
"""
from __future__ import annotations

import argparse
import csv
import errno
import hashlib
import io
import json
import logging
import math
import os
import random
import re
import secrets
import sys
import tempfile
import time
import unicodedata
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, time as datetime_time, timedelta, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Tuple, Union

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError as exc:
    raise RuntimeError("缺少依赖：请安装 openpyxl>=3.1") from exc

HEADERS = ("Task_ID", "Query_Keyword", "Expected_Author", "Scheduled_Time",
           "Status", "Executed_Actions", "Result_Message")
OPTIONAL_HEADERS = ("Expected_Douyin_ID", "Search_Section", "Content_Type", "Comments",
                    "Target_Device_Count")
TERMINAL_STATUSES = ("SUCCESS", "NOT_FOUND", "FAILED")
ACTIONS = ("Like", "Comment", "Favorite", "Share")
REQUIRED_ACTIONS = ("Like", "Comment", "Favorite")
DEFAULTS = {"worksheet": "", "timezone_offset_minutes": 480,
            "file_retry_count": 5, "file_retry_delay_seconds": 0.25,
            "lock_timeout_seconds": 15.0,
            "cooldown_min_seconds": 300, "cooldown_max_seconds": 600,
            "share_probability": 0.5, "sender_accounts": []}


class EngineError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        super().__init__(message)


class TaskStatus(str, Enum):
    PENDING = "PENDING"
    RUNNING = "RUNNING"
    SUCCESS = "SUCCESS"
    NOT_FOUND = "NOT_FOUND"
    FAILED = "FAILED"


class DeviceState(str, Enum):
    READY = "READY"
    BUSY = "BUSY"
    COOLDOWN = "COOLDOWN"
    BLOCKED = "BLOCKED"


class AssertionResult(str, Enum):
    MATCH = "MATCH"
    MISMATCH = "MISMATCH"
    UNCERTAIN = "UNCERTAIN"


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, allow_nan=False,
                      separators=(",", ":"))


def _copy(value: Any) -> Any:
    return json.loads(_json(value))


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _timestamp(value: str) -> datetime:
    result = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if result.tzinfo is None:
        raise ValueError("持久化时间必须包含时区")
    return result.astimezone(timezone.utc)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise EngineError("INVALID_TEXT", "文本字段必须为字符串或数字")
    if isinstance(value, float):
        if not math.isfinite(value):
            raise EngineError("INVALID_TEXT", "不允许 NaN 或无穷数值")
        if value.is_integer():
            value = int(value)
    return str(value).strip()


def _task_id(value: Any) -> str:
    result = _text(value)
    if not result or len(result) > 256:
        raise EngineError("INVALID_TASK_ID", "Task_ID 必须非空且不超过 256 字符")
    return result


def _target_count(value: Any) -> Optional[int]:
    """Blank retains the legacy input contract; explicit counts require allocation."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise EngineError("INVALID_TARGET_COUNT", "Target_Device_Count 必须为正整数")
    try:
        text = _text(value)
    except EngineError as exc:
        raise EngineError("INVALID_TARGET_COUNT", "Target_Device_Count 必须为正整数") from exc
    if not re.fullmatch(r"[1-9]\d*", text):
        raise EngineError("INVALID_TARGET_COUNT", "Target_Device_Count 必须为正整数")
    return int(text)


def _connected_devices(values: Any) -> List[Dict[str, str]]:
    """A fresh actual-device snapshot supplied by ShadowBot, never settings.json."""
    if not isinstance(values, list) or not values:
        raise EngineError("CONNECTED_DEVICES_REQUIRED", "请传入影刀当前实际连接的设备及执行账号列表")
    devices, accounts, result = set(), set(), []
    for value in values:
        if not isinstance(value, dict):
            raise EngineError("INVALID_CONNECTED_DEVICES", "设备项须包含device_id和sender_account_id")
        physical, account = value.get("device_id"), value.get("sender_account_id")
        if any(not isinstance(v, str) or not v.strip() or v != v.strip() or len(v) > 256
               for v in (physical, account)):
            raise EngineError("INVALID_CONNECTED_DEVICES", "设备和账号标识须为非空固定字符串")
        if physical in devices or account in accounts:
            raise EngineError("DUPLICATE_DEVICE_ACCOUNT", "同一批连接设备不能重复设备或执行账号")
        devices.add(physical)
        accounts.add(account)
        result.append({"device_id": physical, "sender_account_id": account})
    return result


def _douyin_id(value: Any) -> str:
    """Strip the UI label; retain identifier case and all significant characters."""
    return re.sub(r"^抖音号\s*[:：]\s*", "", _text(value)).strip()


def _validate_route(search_section: str, content_type: str, columns: Dict[str, int]) -> None:
    if content_type not in ("图文", "视频"):
        raise EngineError("INVALID_CONTENT_TYPE", "Content_Type 必须填写图文或视频")


def _comment_pool(value: Any) -> List[str]:
    if not isinstance(value, str):
        raise EngineError("COMMENTS_EMPTY", "Comments 必须填写至少一条评论")
    comments = list(dict.fromkeys(line.strip() for line in value.splitlines() if line.strip()))
    if not comments:
        raise EngineError("COMMENTS_EMPTY", "Comments 没有有效评论")
    return comments


def _source_key(task: Dict[str, Any]) -> str:
    values = [task.get(name, "") for name in
              ("keyword", "expected_author", "expected_douyin_id", "content_type")]
    return hashlib.sha256(_json(values).encode("utf-8")).hexdigest()


def _atomic_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, name = tempfile.mkstemp(prefix="." + path.name + ".", dir=str(path.parent))
    try:
        with os.fdopen(handle, "wb") as stream:
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, str(path))
        if os.name != "nt":
            descriptor = os.open(str(path.parent), os.O_RDONLY)
            try:
                os.fsync(descriptor)
            finally:
                os.close(descriptor)
    finally:
        if os.path.exists(name):
            os.unlink(name)


@contextmanager
def _project_lock(path: Path, timeout: float) -> Iterator[None]:
    """OS-owned lock, released on process death; never delete the lock file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a+b") as stream:
        stream.seek(0, os.SEEK_END)
        if stream.tell() == 0:
            stream.write(b"\0")
            stream.flush()
        deadline = time.monotonic() + timeout
        while True:
            try:
                stream.seek(0)
                if os.name == "nt":
                    import msvcrt
                    msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(stream.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except OSError as exc:
                if exc.errno not in (errno.EACCES, errno.EAGAIN, errno.EDEADLK):
                    raise
                if time.monotonic() >= deadline:
                    raise EngineError("LOCK_TIMEOUT", "等待项目互斥锁超时") from exc
                time.sleep(0.1)
        try:
            yield
        finally:
            stream.seek(0)
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(stream.fileno(), fcntl.LOCK_UN)


class TaskScheduler:
    """Excel operations; its owner must hold the project lock."""

    def __init__(self, root: Path, settings: Dict[str, Any],
                 logger: logging.Logger) -> None:
        self.path = root / "data" / "task_cases.xlsx"
        self.settings, self.logger = settings, logger
        self.tz = timezone(timedelta(minutes=settings["timezone_offset_minutes"]))

    def _retry(self, operation: Any) -> Any:
        attempts = self.settings["file_retry_count"]
        for attempt in range(attempts):
            try:
                owner_file = self.path.with_name("~$" + self.path.name)
                if owner_file.exists():
                    raise PermissionError("Excel 编辑占用标记存在")
                return operation()
            except FileNotFoundError as exc:
                raise EngineError("EXCEL_MISSING", "找不到用例表：" + str(self.path)) from exc
            except (OSError, zipfile.BadZipFile, InvalidFileException) as exc:
                if attempt + 1 == attempts:
                    raise EngineError("EXCEL_IO_FAILED", "Excel 读写重试耗尽：" + str(exc)) from exc
                self.logger.warning("excel_retry attempt=%s reason=%s", attempt + 1, exc)
                time.sleep(min(2.0, self.settings["file_retry_delay_seconds"] * (2 ** attempt)))

    def read(self) -> Tuple[Any, Any, Dict[str, int], Dict[str, int], str]:
        data = self._retry(self.path.read_bytes)
        try:
            workbook = load_workbook(io.BytesIO(data), data_only=False)
            sheet_name = self.settings["worksheet"]
            sheet = workbook[sheet_name] if sheet_name else workbook.active
            columns: Dict[str, int] = {}
            for cell in sheet[1]:
                name = _text(cell.value)
                if name in HEADERS + OPTIONAL_HEADERS:
                    if name in columns:
                        raise EngineError("EXCEL_SCHEMA_INVALID", "重复字段：" + name)
                    columns[name] = cell.column
            missing = set(HEADERS + ("Content_Type", "Comments")) - set(columns)
            if missing:
                raise EngineError("EXCEL_SCHEMA_INVALID", "缺少字段：" + ",".join(sorted(missing)))
            rows: Dict[str, int] = {}
            for row in range(2, sheet.max_row + 1):
                cells = [sheet.cell(row, columns[name])
                         for name in HEADERS + OPTIONAL_HEADERS if name in columns]
                if all(cell.value is None for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise EngineError("EXCEL_FORMULA_FORBIDDEN", "用例字段不允许公式；行 " + str(row))
                identity = _task_id(sheet.cell(row, columns["Task_ID"]).value)
                if identity in rows:
                    raise EngineError("DUPLICATE_TASK_ID", "重复 Task_ID：" + identity)
                status = _text(sheet.cell(row, columns["Status"]).value).upper()
                if status not in {item.value for item in TaskStatus}:
                    raise EngineError("INVALID_TASK_STATUS", "非法状态；Task_ID=" + identity)
                rows[identity] = row
            return workbook, sheet, columns, rows, hashlib.sha256(data).hexdigest()
        except EngineError:
            raise
        except Exception as exc:
            raise EngineError("EXCEL_FORMAT_INVALID", "无法解析用例表：" + str(exc)) from exc

    def apply_update(self, update: Dict[str, Any]) -> None:
        workbook, sheet, columns, rows, digest = self.read()
        try:
            identity = update["task_id"]
            if identity not in rows:
                raise EngineError("TASK_MISSING", "任务已被移除：" + identity)
            row = rows[identity]
            current = _text(sheet.cell(row, columns["Status"]).value).upper()
            fields = update["fields"]
            if current == fields["Status"]:
                if all(sheet.cell(row, columns[key]).value == value
                       for key, value in fields.items()):
                    return
                if "old_fields" not in update:
                    raise EngineError("TASK_WRITE_CONFLICT", "同状态内容已改变：" + identity)
            if current != update["old_status"]:
                raise EngineError("TASK_WRITE_CONFLICT", "任务状态已改变：" + identity)
            if "old_fields" in update and any(
                    sheet.cell(row, columns[key]).value != value
                    for key, value in update["old_fields"].items()):
                raise EngineError("TASK_WRITE_CONFLICT", "账号结果被外部修改：" + identity)
            for key, value in fields.items():
                sheet.cell(row, columns[key]).value = value
                sheet.cell(row, columns[key]).data_type = "s"
            output = io.BytesIO()
            workbook.save(output)
            payload = output.getvalue()

            def write() -> None:
                if hashlib.sha256(self.path.read_bytes()).hexdigest() != digest:
                    raise EngineError("EXCEL_WRITE_CONFLICT", "Excel 被外部修改，拒绝覆盖")
                _atomic_bytes(self.path, payload)

            self._retry(write)
        finally:
            workbook.close()

    @staticmethod
    def task_from_row(sheet: Any, columns: Dict[str, int], row: int,
                      identity: str) -> Dict[str, Any]:
        content_type = _text(sheet.cell(row, columns["Content_Type"]).value)
        _validate_route(content_type, content_type, columns)
        task = {"task_id": identity,
                "keyword": _text(sheet.cell(row, columns["Query_Keyword"]).value),
                "expected_author": _text(sheet.cell(row, columns["Expected_Author"]).value),
                "expected_douyin_id": (_douyin_id(sheet.cell(row, columns["Expected_Douyin_ID"]).value)
                                       if "Expected_Douyin_ID" in columns else ""),
                "content_type": content_type, "search_section": content_type,
                "comments": _comment_pool(sheet.cell(row, columns["Comments"]).value),
                "target_device_count": (_target_count(sheet.cell(row, columns["Target_Device_Count"]).value)
                                        if "Target_Device_Count" in columns else None)}
        if not task["keyword"] or not task["expected_author"]:
            raise EngineError("INVALID_TASK", "关键词和预期作者必须非空")
        return task

    def due_at(self, identity: str, raw: Any, state: Dict[str, Any],
               now: datetime) -> datetime:
        if isinstance(raw, datetime):
            return (raw.replace(tzinfo=self.tz) if raw.tzinfo is None else raw).astimezone(timezone.utc)
        if isinstance(raw, datetime_time):
            parsed = datetime.combine(now.astimezone().date(), raw)
            return (parsed.replace(tzinfo=self.tz) if parsed.tzinfo is None else parsed).astimezone(timezone.utc)
        text = _text(raw)
        if re.fullmatch(r"\d{2}:\d{2}(?::\d{2})?", text):
            try:
                clock_time = datetime_time.fromisoformat(text)
            except ValueError as exc:
                raise EngineError("INVALID_SCHEDULE", "非法时分秒：" + text) from exc
            return self.due_at(identity, clock_time, state, now)
        relative = re.fullmatch(r"\+?(\d+(?:\.\d+)?)\s*(?:m|min|分钟)?", text, re.I)
        if relative:
            minutes = float(relative.group(1))
            if not math.isfinite(minutes) or minutes > 525600:
                raise EngineError("INVALID_SCHEDULE", "相对分钟数必须在 0～525600 之间")
            cache = state["relative_schedules"].get(identity)
            if not cache or cache["source"] != text:
                cache = {"source": text, "due_at": (now + timedelta(minutes=minutes)).isoformat()}
                state["relative_schedules"][identity] = cache
            return _timestamp(cache["due_at"])
        if not text:
            raise EngineError("INVALID_SCHEDULE", "Scheduled_Time 不允许为空")
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            if "T" not in text and " " not in text:
                raise ValueError("必须提供日期和时间")
            return (parsed.replace(tzinfo=self.tz) if parsed.tzinfo is None else parsed).astimezone(timezone.utc)
        except ValueError as exc:
            raise EngineError("INVALID_SCHEDULE", "不支持的计划时间：" + text) from exc

    def select(self, state: Dict[str, Any], now: datetime,
               device_id: str = "default-device", allocation_only: bool = False) -> Optional[Dict[str, Any]]:
        workbook, sheet, columns, rows, _ = self.read()
        try:
            owners = {device["active_task"]["task_id"]
                      for device in state["devices"].values() if device.get("active_task")}
            projects = state["project_accounts"]
            candidates: List[Tuple[datetime, int, Dict[str, Any]]] = []
            invalid: List[str] = []
            for identity, row in rows.items():
                status = _text(sheet.cell(row, columns["Status"]).value).upper()
                project = projects.get(identity)
                account = project["accounts"].get(device_id) if project else None
                assignment = project.get("assignment") if project else None
                if allocation_only and not assignment:
                    count = sheet.cell(row, columns["Target_Device_Count"]).value if "Target_Device_Count" in columns else None
                    if count is None or count == "":
                        continue
                if status == "RUNNING" and identity not in owners and not assignment:
                    raise EngineError("ORPHAN_RUNNING", "缺少运行归属，需人工核验：" + identity)
                if assignment and device_id not in assignment["selected_accounts"]:
                    continue
                if account is not None and (not assignment or account["status"] != "PENDING"):
                    continue
                if status != "PENDING" and project is None:
                    continue
                try:
                    task = self.task_from_row(sheet, columns, row, identity)
                    task["sender_account_id"] = device_id
                    if project and project["source_key"] != _source_key(task):
                        raise EngineError("TASK_SOURCE_CHANGED", "已执行项目的作品信息被更改；请使用新的Task_ID")
                    if assignment:
                        if task["target_device_count"] != assignment["target_device_count"]:
                            raise EngineError("ASSIGNMENT_CHANGED", "已经分配的设备数量不能更改；请使用新的Task_ID")
                        if account["comment_content"] not in task["comments"]:
                            raise EngineError("ASSIGNMENT_CHANGED", "已经预留的评论不能删除")
                        task.update(assignment_id=assignment["assignment_id"],
                                    physical_device_id=account["physical_device_id"])
                    elif task["target_device_count"] is not None:
                        raise EngineError("ALLOCATION_REQUIRED", "项目" + identity + "需先按影刀实际连接设备分配")
                    configured = self.settings["sender_accounts"]
                    if not assignment and configured and len(task["comments"]) < len(configured):
                        raise EngineError("COMMENTS_INSUFFICIENT", "评论数量少于参与账号数：" + identity)
                    due = self.due_at(identity, sheet.cell(row, columns["Scheduled_Time"]).value, state, now)
                    if due <= now:
                        task["scheduled_time"] = due.isoformat()
                        candidates.append((due, row, task))
                except EngineError as exc:
                    invalid.append(identity + ": " + str(exc))
                    self.logger.error("invalid_task task_id=%s code=%s message=%s", identity, exc.code, exc)
            if candidates:
                return min(candidates, key=lambda item: (item[0], item[1]))[2]
            if invalid:
                raise EngineError("INVALID_TASK_DATA", "; ".join(invalid))
            return None
        finally:
            workbook.close()


class DeviceAndCooldownManager:
    """Device state transitions; JSON persistence is owned by UnifiedGateway."""

    def __init__(self, settings: Dict[str, Any]) -> None:
        minimum = settings.get("cooldown_min_seconds", DEFAULTS["cooldown_min_seconds"])
        maximum = settings.get("cooldown_max_seconds", DEFAULTS["cooldown_max_seconds"])
        for value in (minimum, maximum):
            if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= 2147483647:
                raise EngineError("CONFIG_INVALID", "冷却秒数必须为 0～2147483647 之间的整数")
        if minimum > maximum:
            raise EngineError("CONFIG_INVALID", "cooldown_min_seconds 不能大于 cooldown_max_seconds")
        self.cooldown_min_seconds = minimum
        self.cooldown_max_seconds = maximum

    @staticmethod
    def get(state: Dict[str, Any], device_id: str, now: datetime) -> Dict[str, Any]:
        device = state["devices"].setdefault(device_id, {
            "state": "READY", "active_task": None, "cooldown_until": None,
            "last_result": None, "block_info": None})
        if device["state"] == "COOLDOWN" and now >= _timestamp(device["cooldown_until"]):
            device["state"], device["cooldown_until"] = "READY", None
        return device

    @staticmethod
    def claim(device: Dict[str, Any], task: Dict[str, Any], now: datetime) -> Dict[str, Any]:
        if device["state"] != "READY" or device.get("active_task"):
            raise EngineError("DEVICE_UNAVAILABLE", "设备未就绪")
        active = dict(task, run_token=uuid.uuid4().hex, claimed_at=now.isoformat(), strategy=None)
        device.update(state="BUSY", active_task=active, cooldown_until=None)
        return active

    def cool(self, device: Dict[str, Any], now: datetime) -> int:
        minimum, maximum = self.cooldown_min_seconds, self.cooldown_max_seconds
        seconds = secrets.randbelow(maximum - minimum + 1) + minimum
        device.update(state="COOLDOWN", active_task=None,
                      cooldown_until=(now + timedelta(seconds=seconds)).isoformat())
        return seconds


class InteractionStrategyEngine:
    def __init__(self, share_probability: float = 0.5) -> None:
        self.share_probability = share_probability

    def generate(self, seed: Optional[Union[int, str]] = None,
                 comment_content: str = "") -> Dict[str, Any]:
        if not comment_content:
            raise EngineError("COMMENTS_EMPTY", "当前账号尚未分配项目评论")
        if seed is None:
            seed = secrets.randbits(64)
        if isinstance(seed, bool) or not re.fullmatch(r"-?\d+", str(seed)):
            raise EngineError("INVALID_SEED", "随机种子必须为整数或整数字符串")
        rng = random.Random(int(seed))
        order = list(REQUIRED_ACTIONS)
        share_selected = rng.random() < self.share_probability
        rng.shuffle(order)
        if share_selected:
            order.append("Share")
        return {"policy_version": 2, "seed": str(seed), "actions": order[:],
                "execution_order": order, "comment_content": comment_content,
                "share_selected": "Share" in order, "share_destination": "friend",
                "delays": {"initial_browse": round(rng.uniform(2.5, 4.5), 3),
                           "action_gaps": [round(rng.uniform(1.8, 3.2), 3)
                                           for _ in range(len(order) - 1)]}}


class AuthorAssertionHelper:
    def __init__(self, whitelist_path: Path) -> None:
        self.path = whitelist_path

    @staticmethod
    def normalize(value: Any) -> str:
        text = unicodedata.normalize("NFKC", _text(value))
        text = text.replace("\ufeff", "").replace("\u200b", "")
        return re.sub(r"\s+", " ", text).strip()

    def _names(self) -> Dict[str, List[str]]:
        if not self.path.exists():
            return {}
        try:
            data = json.loads(self.path.read_text(encoding="utf-8-sig"))
            if not isinstance(data, dict):
                raise ValueError("白名单必须是 JSON 对象")
            entries = data.get("authors")
            if entries is None:
                entries = [{"uid": uid, "nicknames": [name]} for name, uid in data.items()]
            if not isinstance(entries, list):
                raise ValueError("authors 必须为数组")
            names: Dict[str, List[str]] = {}
            for entry in entries:
                uid = self.normalize(entry["uid"])
                nicknames = entry["nicknames"]
                if not uid or not isinstance(nicknames, list):
                    raise ValueError("每项须包含非空 uid 和 nicknames 数组")
                for nickname in nicknames:
                    name = self.normalize(nickname)
                    if not name:
                        raise ValueError("白名单昵称不允许为空")
                    bucket = names.setdefault(name, [])
                    if uid not in bucket:
                        bucket.append(uid)
            return names
        except (OSError, ValueError, KeyError, TypeError, EngineError) as exc:
            raise EngineError("WHITELIST_INVALID", "无法解析作者白名单：" + str(exc)) from exc

    def assert_author(self, expected: Any, actual: Any) -> str:
        # Explicit Douyin ID checks must not consult the internal-UID whitelist
        # or silently fall back to matching nicknames when an ID is missing.
        if isinstance(expected, dict) and "douyin_id" in expected:
            def unpack_douyin(value: Any) -> Tuple[str, bool]:
                if isinstance(value, dict):
                    complete = value.get("complete", True)
                    if not isinstance(complete, bool):
                        raise EngineError("INVALID_AUTHOR", "complete 必须是布尔值")
                    source = _text(value.get("source", "ui")).lower()
                    return _douyin_id(value.get("douyin_id")), complete and source in ("ui", "verified")
                return _douyin_id(value), True

            expected_id, expected_ok = unpack_douyin(expected)
            actual_id, actual_ok = unpack_douyin(actual)
            for identity, verified in ((expected_id, expected_ok), (actual_id, actual_ok)):
                if (not verified or not identity or "..." in identity or "…" in identity
                        or re.search(r"\s|[:：]", identity)):
                    return "UNCERTAIN"
            return "MATCH" if expected_id == actual_id else "MISMATCH"

        def unpack(value: Any) -> Tuple[str, str, bool]:
            if isinstance(value, dict):
                complete = value.get("complete", True)
                if not isinstance(complete, bool):
                    raise EngineError("INVALID_AUTHOR", "complete 必须是布尔值")
                source = _text(value.get("source", "ui")).lower()
                verified = complete and source in ("ui", "verified")
                return self.normalize(value.get("uid")), self.normalize(value.get("nickname")), verified
            name = self.normalize(value)
            if name.lower().startswith("uid:"):
                return self.normalize(name[4:]), "", True
            return "", name, True

        expected_uid, expected_name, expected_complete = unpack(expected)
        actual_uid, actual_name, actual_complete = unpack(actual)
        if not expected_complete or not actual_complete:
            return "UNCERTAIN"
        if expected_uid:
            if expected_uid.endswith(("...", "…")) or actual_uid.endswith(("...", "…")):
                return "UNCERTAIN"
            return ("MATCH" if expected_uid == actual_uid else "MISMATCH") if actual_uid else "UNCERTAIN"
        if not expected_name or expected_name.endswith(("...", "…")):
            return "UNCERTAIN"
        if not actual_uid:
            if not actual_name or actual_name.endswith(("...", "…")):
                return "UNCERTAIN"
            return "MATCH" if expected_name == actual_name else "MISMATCH"
        identities = self._names().get(expected_name, [])
        if len(identities) > 1:
            return "UNCERTAIN"
        if identities:
            if actual_uid.endswith(("...", "…")):
                return "UNCERTAIN"
            return ("MATCH" if identities[0] == actual_uid else "MISMATCH") if actual_uid else "UNCERTAIN"
        if not actual_name or actual_name.endswith(("...", "…")):
            return "UNCERTAIN"
        return "MATCH" if expected_name == actual_name else "MISMATCH"


class UnifiedGateway:
    """Public gateway. Returned values contain only JSON-compatible types."""

    def __init__(self, root: Optional[str] = None,
                 device_id: str = "default-device") -> None:
        default_root = Path(__file__).resolve().parent
        if default_root.name == "scripts":
            default_root = default_root.parent
        self.root = Path(root or os.environ.get("PHONEWORKBENCH_ROOT") or
                         str(default_root)).expanduser().resolve()
        self.device_id = _text(device_id)
        if not self.device_id:
            raise EngineError("INVALID_DEVICE_ID", "device_id 必须非空")
        for folder in ("config", "data", "logs"):
            (self.root / folder).mkdir(parents=True, exist_ok=True)
        logger_name = "PhoneWorkbench." + hashlib.sha256(str(self.root).encode()).hexdigest()[:16]
        self.logger = logging.getLogger(logger_name)
        self.logger.setLevel(logging.INFO)
        self.logger.propagate = False
        if not self.logger.handlers:
            handler = logging.FileHandler(str(self.root / "logs" / "execution.log"), encoding="utf-8")
            handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s pid=%(process)d %(message)s"))
            self.logger.addHandler(handler)
        self.settings = dict(DEFAULTS)
        config = self.root / "config" / "settings.json"
        if config.exists():
            try:
                overrides = json.loads(config.read_text(encoding="utf-8-sig"))
                if not isinstance(overrides, dict):
                    raise ValueError("配置须为 JSON 对象")
                self.settings.update(overrides)
            except (OSError, ValueError) as exc:
                raise EngineError("CONFIG_INVALID", "配置读取失败：" + str(exc)) from exc
        for key, low, high in (("timezone_offset_minutes", -840, 840),
                               ("file_retry_count", 1, 8)):
            value = self.settings[key]
            if isinstance(value, bool) or not isinstance(value, int) or not low <= value <= high:
                raise EngineError("CONFIG_INVALID", "配置整数超出范围：" + key)
        for key, low, high in (("file_retry_delay_seconds", 0.01, 2.0),
                               ("lock_timeout_seconds", 0.1, 60.0)):
            value = self.settings[key]
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not low <= value <= high:
                raise EngineError("CONFIG_INVALID", "配置数值超出范围：" + key)
        if not isinstance(self.settings["worksheet"], str):
            raise EngineError("CONFIG_INVALID", "worksheet 必须为字符串")
        probability = self.settings["share_probability"]
        if (isinstance(probability, bool) or not isinstance(probability, (int, float))
                or not math.isfinite(probability) or not 0 <= probability <= 1):
            raise EngineError("CONFIG_INVALID", "share_probability 必须在0到1之间")
        accounts = self.settings["sender_accounts"]
        if (not isinstance(accounts, list)
                or any(not isinstance(value, str) or not value.strip() or value != value.strip()
                       for value in accounts) or len(accounts) != len(set(accounts))):
            raise EngineError("CONFIG_INVALID", "sender_accounts 必须为不重复的账号标识列表")
        if accounts and self.device_id not in accounts:
            raise EngineError("ACCOUNT_NOT_CONFIGURED", "当前账号不在sender_accounts中")
        self.state_path = self.root / "data" / "runtime_state.json"
        self.lock_path = self.root / "data" / ".core_runner.lock"
        self.scheduler = TaskScheduler(self.root, self.settings, self.logger)
        self.devices = DeviceAndCooldownManager(self.settings)
        self.strategy = InteractionStrategyEngine(probability)
        self.authors = AuthorAssertionHelper(self.root / "data" / "author_whitelist.json")

    def _load_state(self) -> Dict[str, Any]:
        if not self.state_path.exists():
            return {"schema_version": 1, "devices": {}, "relative_schedules": {},
                    "pending_transaction": None, "project_accounts": {}}
        try:
            state = json.loads(self.state_path.read_text(encoding="utf-8"))
            if state["schema_version"] != 1:
                raise ValueError("不兼容的状态文件版本")
            if not isinstance(state["devices"], dict) or not isinstance(state["relative_schedules"], dict):
                raise ValueError("设备与相对时间缓存须为对象")
            pending = state["pending_transaction"]
            if pending is not None:
                if not isinstance(pending, dict) or pending["old_status"] not in {item.value for item in TaskStatus}:
                    raise ValueError("事务记录无效")
                _task_id(pending["task_id"])
                if set(pending["fields"]) != {"Status", "Executed_Actions", "Result_Message"}:
                    raise ValueError("事务字段无效")
                if pending["fields"]["Status"] not in {item.value for item in TaskStatus}:
                    raise ValueError("事务目标状态无效")
            for cache in state["relative_schedules"].values():
                if not isinstance(cache["source"], str):
                    raise ValueError("相对时间缓存无效")
                _timestamp(cache["due_at"])
            for device in state["devices"].values():
                if device["state"] not in {item.value for item in DeviceState}:
                    raise ValueError("非法设备状态")
                active = device["active_task"]
                if active is not None:
                    _task_id(active["task_id"])
                    if not isinstance(active["run_token"], str) or not active["run_token"]:
                        raise ValueError("运行令牌无效")
                    _timestamp(active["claimed_at"])
                    if device["state"] not in ("BUSY", "BLOCKED"):
                        raise ValueError("活动任务的设备状态不一致")
                elif device["state"] == "BUSY":
                    raise ValueError("BUSY 设备缺少任务")
                if device["state"] == "COOLDOWN":
                    _timestamp(device["cooldown_until"])
                if device["state"] == "BLOCKED" and not isinstance(device["block_info"], dict):
                    raise ValueError("BLOCKED 设备缺少熔断信息")
            projects = state.setdefault("project_accounts", {})
            if not isinstance(state.setdefault("exports_dirty", False), bool):
                raise ValueError("输出恢复标记无效")
            if not isinstance(projects, dict):
                raise ValueError("项目账号进度须为对象")
            for identity, project in projects.items():
                _task_id(identity)
                if not isinstance(project["source_key"], str) or not isinstance(project["accounts"], dict):
                    raise ValueError("项目账号进度无效")
                comments = []
                assignment = project.get("assignment")
                if assignment is not None:
                    roster = _connected_devices(assignment["connected_devices"])
                    selected = assignment["selected_accounts"]
                    count = _target_count(assignment["target_device_count"])
                    if (not isinstance(selected, list) or len(selected) != count
                            or len(set(selected)) != count or set(selected) != set(project["accounts"])
                            or not set(selected).issubset({v["sender_account_id"] for v in roster})
                            or not isinstance(assignment["assignment_id"], str) or not assignment["assignment_id"]):
                        raise ValueError("固定设备分配记录无效")
                    _timestamp(assignment["allocated_at"])
                    bindings = {v["sender_account_id"]: v["device_id"] for v in roster}
                    for account, entry in project["accounts"].items():
                        if entry.get("physical_device_id") != bindings[account]:
                            raise ValueError("分配账号与物理设备不一致")
                        if entry["status"] == "PENDING" and (entry.get("run_token") or entry.get("result")):
                            raise ValueError("待领取分配已有运行结果")
                for account, entry in project["accounts"].items():
                    if not isinstance(account, str) or not account:
                        raise ValueError("发送账号标识无效")
                    if entry["status"] not in {item.value for item in TaskStatus}:
                        raise ValueError("账号任务状态无效")
                    if not isinstance(entry.get("comment_content", ""), str):
                        raise ValueError("分配评论无效")
                    if entry.get("comment_content"):
                        comments.append(entry["comment_content"])
                    if entry["status"] == "RUNNING":
                        active = state["devices"].get(account, {}).get("active_task")
                        if (not active or active["task_id"] != identity
                                or active["run_token"] != entry["run_token"]
                                or active.get("comment_content") != entry["comment_content"]):
                            raise ValueError("活动账号与项目分配不一致")
                        plan = active.get("strategy")
                        if plan and plan.get("policy_version") == 2 and plan["comment_content"] != entry["comment_content"]:
                            raise ValueError("互动计划与分配评论不一致")
                        if assignment and (active.get("assignment_id") != assignment["assignment_id"]
                                           or active.get("physical_device_id") != entry["physical_device_id"]):
                            raise ValueError("活动任务与固定设备分配不一致")
                        if "action_results" in active:
                            self._actions(active["action_results"])
                    if assignment and entry["status"] in TERMINAL_STATUSES:
                        result = entry.get("result")
                        if not result or result["status"] != entry["status"]:
                            raise ValueError("结束分配缺少结果")
                        self._actions(result["executed_actions"])
                if len(comments) != len(set(comments)):
                    raise ValueError("同一项目存在重复分配评论")
            return state
        except (OSError, ValueError, KeyError, TypeError, AttributeError, EngineError) as exc:
            raise EngineError("STATE_CORRUPT", "状态文件损坏，拒绝重建或解锁：" + str(exc)) from exc

    def _save_state(self, state: Dict[str, Any]) -> None:
        try:
            _atomic_bytes(self.state_path, (_json(state) + "\n").encode("utf-8"))
        except (OSError, ValueError, TypeError) as exc:
            raise EngineError("STATE_WRITE_FAILED", "状态文件写入失败：" + str(exc)) from exc

    def _commit(self, state: Dict[str, Any], update: Dict[str, Any]) -> None:
        """Write-ahead journal: persist device gate BEFORE changing Excel."""
        if state["pending_transaction"] is not None:
            raise EngineError("TRANSACTION_PENDING", "上次事务尚未完成")
        state["pending_transaction"] = update
        state["exports_dirty"] = True
        self._save_state(state)
        self._replay(state)
        self._flush_exports(state)

    def _replay(self, state: Dict[str, Any]) -> None:
        update = state["pending_transaction"]
        if update is not None:
            try:
                self.scheduler.apply_update(update)
            except Exception as exc:
                self.logger.exception("transaction_pending task_id=%s", update["task_id"])
                raise EngineError("WRITE_PENDING", "事务已持久化，设备保持锁定；修复 Excel 后重试：" + str(exc)) from exc
            state["pending_transaction"] = None
            self._save_state(state)
            self.logger.info("transaction_committed task_id=%s status=%s",
                             update["task_id"], update["fields"]["Status"])

    def _recover(self, state: Dict[str, Any]) -> None:
        self._replay(state)
        # A security block can preempt a claim even while Excel is occupied.
        for identity, device in state["devices"].items():
            if device["state"] == "BLOCKED" and device.get("active_task"):
                info = device["block_info"]
                self._finish(state, identity, device, "FAILED", info["executed_actions"],
                             "SECURITY_CHALLENGE_DETECTED", info["log_message"], True)
        self._flush_exports(state)

    def _register_active(self, state: Dict[str, Any], account_id: str,
                         active: Dict[str, Any]) -> Dict[str, Any]:
        """Reserve under the existing project lock; never reassign another account's text."""
        project = state["project_accounts"].setdefault(active["task_id"], {
            "source_key": _source_key(active), "accounts": {}})
        entry = project["accounts"].get(account_id)
        if entry and entry.get("run_token") not in (None, active["run_token"]):
            raise EngineError("STATE_CONFLICT", "账号任务令牌与分配记录不一致")
        if active.get("strategy") is not None and not entry:
            raise EngineError("LEGACY_STRATEGY", "已有旧版互动计划，请先核对执行进度")
        reserved = {value["comment_content"] for account, value in project["accounts"].items()
                    if account != account_id and value.get("comment_content")}
        pool = active.get("comments", [])
        comment = entry.get("comment_content", "") if entry else active.get("comment_content", "")
        if comment not in pool:
            if project.get("assignment"):
                raise EngineError("ASSIGNMENT_CHANGED", "固定分配已预留评论，不能自动更换")
            available = [value for value in pool if value not in reserved]
            if not available:
                raise EngineError("COMMENTS_EXHAUSTED", "没有未分配的评论，请补充Comments")
            comment = secrets.choice(available)
        if comment in reserved:
            raise EngineError("STATE_CONFLICT", "评论已属于另一个账号")
        active.update(comment_content=comment, sender_account_id=account_id)
        if entry is None:
            entry = {"status": "RUNNING", "run_token": active["run_token"],
                     "claimed_at": active["claimed_at"], "comment_sent": False, "result": None}
            project["accounts"][account_id] = entry
        elif entry["status"] == "PENDING":
            entry.update(status="RUNNING", run_token=active["run_token"],
                         claimed_at=active["claimed_at"])
        entry["comment_content"] = comment
        if active.get("physical_device_id"):
            entry["physical_device_id"] = active["physical_device_id"]
        return entry

    def _set_workspace(self, active: Dict[str, Any]) -> None:
        """Separate OCR/screenshot files for every account and claim."""
        account_key = hashlib.sha256(self.device_id.encode("utf-8")).hexdigest()[:24]
        token_key = hashlib.sha256(active["run_token"].encode("utf-8")).hexdigest()
        path = self.root / "workspaces" / account_key / token_key
        path.mkdir(parents=True, exist_ok=True)
        active["workspace_path"] = str(path)

    def _project_update(self, state: Dict[str, Any], task_id: str) -> Dict[str, Any]:
        entries = state["project_accounts"][task_id]["accounts"]
        statuses = [entry["status"] for entry in entries.values()]
        assignment = state["project_accounts"][task_id].get("assignment")
        configured = assignment["selected_accounts"] if assignment else self.settings["sender_accounts"]
        waiting = [account for account in configured if account not in entries]
        status = (self._fixed_status(statuses) if assignment else
                  "RUNNING" if "RUNNING" in statuses else
                  "PENDING" if waiting else
                  "FAILED" if "FAILED" in statuses else
                  "SUCCESS" if "SUCCESS" in statuses else "NOT_FOUND")
        actions = {}
        summaries = {}
        for account, entry in entries.items():
            result = entry.get("result") or {}
            actions[account] = [{"action": value["action"], "status": value["status"]}
                               for value in result.get("executed_actions", [])]
            summaries[account] = {"status": entry["status"],
                                  "comment_sent": entry.get("comment_sent", False),
                                  "finished_at": result.get("finished_at"),
                                  "error_code": result.get("error_code", "")}
        for account in waiting:
            actions[account] = []
            summaries[account] = {"status": "PENDING", "comment_sent": False,
                                  "finished_at": None, "error_code": ""}
        summary = {"scope": "selected_sender_accounts" if assignment else
                            "configured_sender_accounts" if configured else "known_sender_accounts",
                   "accounts": summaries}
        if assignment:
            summary.update(assignment_id=assignment["assignment_id"],
                           target_device_count=assignment["target_device_count"],
                           finished_device_count=sum(v in TERMINAL_STATUSES for v in statuses))
        fields = {"Status": status, "Executed_Actions": _json({"accounts": actions}),
                  "Result_Message": _json(summary)}
        if any(len(value) > 32767 for value in fields.values()):
            raise EngineError("RESULT_TOO_LONG", "账号汇总超过Excel单元格上限；完整结果保留在状态文件")
        workbook, sheet, columns, rows, _ = self.scheduler.read()
        try:
            if task_id not in rows:
                raise EngineError("TASK_MISSING", "任务已被移除：" + task_id)
            row = rows[task_id]
            old_fields = {name: sheet.cell(row, columns[name]).value for name in fields}
            return {"task_id": task_id, "old_status": _text(old_fields["Status"]).upper(),
                    "old_fields": old_fields, "fields": fields}
        finally:
            workbook.close()

    @staticmethod
    def _fixed_status(statuses: List[str]) -> str:
        if any(value not in TERMINAL_STATUSES for value in statuses):
            return "PENDING" if all(value == "PENDING" for value in statuses) else "RUNNING"
        if all(value == "SUCCESS" for value in statuses):
            return "SUCCESS"
        return "NOT_FOUND" if all(value == "NOT_FOUND" for value in statuses) else "FAILED"

    def allocate_task(self, task_id: Union[str, int],
                      connected_devices: List[Dict[str, str]]) -> Dict[str, Any]:
        """Call from the ShadowBot controller with a current connection snapshot.

        This method never enumerates phones and never treats sender_accounts as
        connected phones. Selection and comment reservation commit exactly once.
        """
        identity = _task_id(task_id)
        roster = _connected_devices(connected_devices)
        configured = self.settings["sender_accounts"]
        if configured and any(v["sender_account_id"] not in configured for v in roster):
            raise EngineError("ACCOUNT_NOT_CONFIGURED", "连接名单含未登记的执行账号")
        with self._session() as state:
            return self._allocate_task_locked(state, identity, roster)

    def _allocate_task_locked(self, state, identity, roster):
        # LAN_TOOL_INTAKE_V1
        configured = self.settings["sender_accounts"]
        if configured and any(v["sender_account_id"] not in configured for v in roster):
            raise EngineError("ACCOUNT_NOT_CONFIGURED", "连接名单含未登记的执行账号")
        workbook, sheet, columns, rows, _ = self.scheduler.read()
        try:
            if identity not in rows:
                raise EngineError("TASK_MISSING", "找不到项目：" + identity)
            row = rows[identity]
            task = self.scheduler.task_from_row(sheet, columns, row, identity)
            count = task["target_device_count"]
            if count is None:
                raise EngineError("TARGET_COUNT_REQUIRED", "请填写Target_Device_Count执行设备数量")
            status = _text(sheet.cell(row, columns["Status"]).value).upper()
            existing = state["project_accounts"].get(identity)
            if existing:
                assignment = existing.get("assignment")
                if not assignment:
                    raise EngineError("LEGACY_PROJECT_REVIEW_REQUIRED", "旧项目已有进度，不能重新随机分配；新分配使用新的Task_ID")
                if (_source_key(task) != existing["source_key"]
                        or count != assignment["target_device_count"]
                        or any(v["comment_content"] not in task["comments"] for v in existing["accounts"].values())):
                    raise EngineError("ASSIGNMENT_CHANGED", "已分配的作品、设备数量或预留评论被更改")
                return self._project_progress(state, identity)
            if status != "PENDING":
                raise EngineError("TASK_NOT_PENDING", "仅允许给新的PENDING项目分配设备")
            if count > len(roster):
                raise EngineError("TARGET_COUNT_EXCEEDS_CONNECTED", "执行设备数量大于当前实际连接设备数量")
            eligible = [v for v in roster if state["devices"].get(v["sender_account_id"], {}).get("state") != "BLOCKED"]
            if count > len(eligible):
                raise EngineError("ELIGIBLE_DEVICES_INSUFFICIENT", "可执行设备不足；被锁定设备不能参与分配")
            for value in roster:
                active = state["devices"].get(value["sender_account_id"], {}).get("active_task")
                if active and active.get("physical_device_id") not in (None, value["device_id"]):
                    raise EngineError("ACCOUNT_DEVICE_CONFLICT", "账号已有任务，不能在执行中更换物理设备")
            if count > len(task["comments"]):
                raise EngineError("COMMENTS_INSUFFICIENT", "去重后的候选评论数量少于本次执行设备数量")
            # Validate the schedule before saving a permanent assignment.
            self.scheduler.due_at(identity, sheet.cell(row, columns["Scheduled_Time"]).value, state, _utcnow())
        finally:
            workbook.close()
        selected = secrets.SystemRandom().sample(eligible, count)
        comments = secrets.SystemRandom().sample(task["comments"], count)
        assignment = {"assignment_id": uuid.uuid4().hex, "allocated_at": _utcnow().isoformat(),
                      "target_device_count": count, "connected_devices": roster,
                      "selected_accounts": [v["sender_account_id"] for v in selected]}
        state["project_accounts"][identity] = {
            "source_key": _source_key(task), "assignment": assignment,
            "accounts": {value["sender_account_id"]: {
                "physical_device_id": value["device_id"], "status": "PENDING",
                "comment_content": comment, "comment_sent": False, "run_token": None,
                "claimed_at": None, "result": None}
                for value, comment in zip(selected, comments)}}
        self._commit(state, self._project_update(state, identity))
        return self._project_progress(state, identity)

    def _project_progress(self, state: Dict[str, Any], identity: str) -> Dict[str, Any]:
        project = state["project_accounts"].get(identity)
        if project is None:
            raise EngineError("PROJECT_NOT_ALLOCATED", "项目尚无分配记录：" + identity)
        assignment = project.get("assignment")
        accounts = (assignment["selected_accounts"] if assignment else
                    list(dict.fromkeys(self.settings["sender_accounts"] + list(project["accounts"]))))
        details = []
        for account in accounts:
            entry = project["accounts"].get(account, {})
            result = entry.get("result")
            active = state["devices"].get(account, {}).get("active_task")
            active = active if active and active["task_id"] == identity else None
            actions = result["executed_actions"] if result else active.get("action_results", []) if active else []
            passed = {v["action"] for v in actions if v["status"] == "PASSED"}
            plan = active.get("strategy") if active else entry.get("plan")
            details.append({"sender_account_id": account,
                            "physical_device_id": entry.get("physical_device_id"),
                            "status": entry.get("status", "PENDING"),
                            "run_token": entry.get("run_token"),
                            "claimed_at": entry.get("claimed_at"),
                            "finished_at": result.get("finished_at") if result else None,
                            "last_progress_at": (result.get("finished_at") if result else
                                                 active.get("last_progress_at", active.get("claimed_at")) if active else None),
                            "resume_requires_review": bool(active and active.get("strategy")),
                            "comment_sent": entry.get("comment_sent", False),
                            "executed_actions": _copy(actions),
                            "planned_actions": plan.get("execution_order", []) if plan else [],
                            "counts": {name: int(name in passed) for name in ACTIONS},
                            "error_code": result.get("error_code", "") if result else "",
                            "log_message": result.get("log_message", "") if result else ""})
        statuses = [v["status"] for v in details]
        finished = sum(value in TERMINAL_STATUSES for value in statuses)
        return {"schema_version": 1, "task_id": identity, "source_key": project["source_key"],
                "allocation_mode": "fixed_random" if assignment else "legacy",
                "assignment_id": assignment["assignment_id"] if assignment else None,
                "allocated_at": assignment["allocated_at"] if assignment else None,
                "target_device_count": len(accounts), "assigned_device_count": len(accounts),
                "finished_device_count": finished,
                "pending_device_count": statuses.count("PENDING"),
                "running_device_count": statuses.count("RUNNING"),
                "success_device_count": statuses.count("SUCCESS"),
                "failed_device_count": statuses.count("FAILED"),
                "not_found_device_count": statuses.count("NOT_FOUND"),
                "completed": bool(accounts) and finished == len(accounts),
                "status": self._fixed_status(statuses) if accounts else "PENDING",
                "counts": {name: sum(v["counts"][name] for v in details) for name in ACTIONS},
                "accounts": details}

    def get_project_progress(self, task_id: Union[str, int]) -> Dict[str, Any]:
        """Stable JSON contract for later frontend/table integrations."""
        with self._session() as state:
            return _copy(self._project_progress(state, _task_id(task_id)))

    def _flush_exports(self, state: Dict[str, Any]) -> None:
        if not state.get("exports_dirty"):
            return
        accounts = {account for project in state["project_accounts"].values()
                    for account, entry in project["accounts"].items() if entry.get("result") is not None}
        for account in sorted(accounts):
            self._write_execution_data(state, account)
        output = {"schema_version": 1, "projects": [self._project_progress(state, identity)
                                                   for identity in state["project_accounts"]]}
        path = self.root / "execution_data" / "project_progress.json"
        try:
            _atomic_bytes(path, (_json(output) + "\n").encode("utf-8"))
        except OSError as exc:
            raise EngineError("EXECUTION_DATA_WRITE_FAILED", "项目进度输出失败：" + str(path)) from exc
        state["exports_dirty"] = False
        self._save_state(state)
        try:
            import lan_bridge
            lan_bridge.snapshot(self, state)
        except ImportError:
            pass

    def save_action_results(self, task_id: Union[str, int], run_token: str,
                            executed_actions: List[Any]) -> List[Dict[str, str]]:
        """Save confirmations after each action; never retry mobile actions here."""
        identity, actions = _task_id(task_id), self._actions(executed_actions)
        with self._session() as state:
            device = self.devices.get(state, self.device_id, _utcnow())
            active = self._owned(device, identity, run_token)
            plan = active.get("strategy")
            if not plan or [v["action"] for v in actions] != plan["execution_order"][:len(actions)]:
                raise EngineError("ACTION_ORDER_CONFLICT", "动作进度必须符合已保存计划的顺序")
            saved = active.get("action_results", [])
            if len(actions) < len(saved) or actions[:len(saved)] != saved:
                raise EngineError("ACTION_RESULT_CONFLICT", "已保存动作结果不能覆盖或减少")
            if actions == saved:
                return _copy(actions)
            active["action_results"] = actions
            active["last_progress_at"] = _utcnow().isoformat()
            entry = state["project_accounts"][identity]["accounts"][self.device_id]
            entry["comment_sent"] = any(v["action"] == "Comment" and v["status"] == "PASSED" for v in actions)
            state["exports_dirty"] = True
            self._save_state(state)
            self._flush_exports(state)
            return _copy(actions)

    def close_pending_assignment(self, task_id: Union[str, int], sender_account_id: str,
                                 operator: str, reason: str) -> Dict[str, Any]:
        """Explicit operator closure only; a running/uncertain task must be reviewed."""
        if any(not isinstance(v, str) or not v.strip() for v in (operator, reason)):
            raise EngineError("MANUAL_RESET_REQUIRED", "结束待领取分配须提供操作人和原因")
        identity = _task_id(task_id)
        with self._session() as state:
            project = state["project_accounts"].get(identity)
            entry = project["accounts"].get(sender_account_id) if project else None
            if not project or not project.get("assignment") or not entry or entry["status"] != "PENDING":
                raise EngineError("ASSIGNMENT_NOT_PENDING", "仅允许明确结束尚未领取的分配")
            now = _utcnow().isoformat()
            entry.update(status="FAILED", result={
                "task_id": identity, "assignment_id": project["assignment"]["assignment_id"],
                "sender_account_id": sender_account_id, "device_id": sender_account_id,
                "physical_device_id": entry["physical_device_id"], "run_token": None,
                "status": "FAILED", "executed_actions": [], "finished_at": now,
                "error_code": "ASSIGNMENT_CLOSED_BY_OPERATOR", "log_message": reason,
                "operator": operator})
            self._commit(state, self._project_update(state, identity))
            return self._project_progress(state, identity)

    def adopt_legacy_task(self, legacy_device_id: str = "default-device") -> bool:
        """Explicitly rename a pre-interaction legacy slot to the selected sender account."""
        if legacy_device_id == self.device_id:
            return False
        with self._session() as state:
            legacy = state["devices"].get(legacy_device_id)
            if not legacy or not legacy.get("active_task"):
                return False
            if legacy["state"] != "BUSY" or legacy["active_task"].get("strategy") is not None:
                raise EngineError("LEGACY_TASK_REVIEW_REQUIRED", "旧任务已生成互动策略或被锁定，请先核对，不能自动转交")
            target = state["devices"].get(self.device_id)
            if target and (target["state"] != "READY" or target.get("active_task") or target.get("last_result")):
                raise EngineError("STATE_CONFLICT", "目标账号已有执行记录，不能转交旧任务")
            if state["project_accounts"].get(legacy["active_task"]["task_id"]):
                raise EngineError("STATE_CONFLICT", "旧任务已有账号分配，不能转交")
            state["devices"][self.device_id] = state["devices"].pop(legacy_device_id)
            state["devices"][self.device_id]["active_task"]["sender_account_id"] = self.device_id
            self._save_state(state)
            return True

    @contextmanager
    def _session(self, recover: bool = True) -> Iterator[Dict[str, Any]]:
        try:
            with _project_lock(self.lock_path, self.settings["lock_timeout_seconds"]):
                state = self._load_state()
                if state.get("web_cleanup_transaction"):
                    import lan_bridge
                    lan_bridge.recover_cleanup(self, state)
                if recover:
                    self._recover(state)
                yield state
        except EngineError as exc:
            self.logger.error("gateway_error device_id=%s code=%s message=%s", self.device_id, exc.code, exc)
            raise
        except Exception as exc:
            self.logger.exception("gateway_error device_id=%s", self.device_id)
            raise EngineError("ENGINE_ERROR", str(exc)) from exc

    @staticmethod
    def _actions(values: Any) -> List[Dict[str, str]]:
        if not isinstance(values, (list, tuple)):
            raise EngineError("INVALID_ACTIONS", "executed_actions 必须为数组")
        results: List[Dict[str, str]] = []
        seen = set()
        for value in values:
            item = {"action": value, "status": "PASSED"} if isinstance(value, str) else value
            if not isinstance(item, dict):
                raise EngineError("INVALID_ACTIONS", "动作项必须为名称或结果对象")
            name, status = item.get("action"), item.get("status")
            if name not in ACTIONS or name in seen:
                raise EngineError("INVALID_ACTIONS", "动作非法或重复")
            if status not in ("PASSED", "FAILED", "UNCERTAIN", "SKIPPED_ALREADY_ACTIVE"):
                raise EngineError("INVALID_ACTIONS", "动作结果状态非法")
            if status == "SKIPPED_ALREADY_ACTIVE" and name not in ("Like", "Favorite"):
                raise EngineError("INVALID_ACTIONS", "只有点赞和收藏支持已激活跳过")
            message = item.get("message", "")
            if not isinstance(message, str) or len(message) > 2000:
                raise EngineError("INVALID_ACTIONS", "动作 message 必须为不超过 2000 字符的字符串")
            results.append({"action": name, "status": status, "message": message})
            seen.add(name)
        return results

    @staticmethod
    def _owned(device: Dict[str, Any], identity: str,
               run_token: Optional[str]) -> Dict[str, Any]:
        if not run_token:
            raise EngineError("RUN_TOKEN_REQUIRED", "必须传入领取任务时返回的 run_token")
        active = device.get("active_task")
        if not active or active["task_id"] != identity or active["run_token"] != run_token:
            raise EngineError("TASK_NOT_OWNED", "任务或运行令牌不属于当前设备")
        return active

    def _write_execution_data(self, state: Dict[str, Any], sender_account_id: str) -> Path:
        """Export this account's recorded results; one row per task, no totals."""
        filename_id = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", sender_account_id).strip(" .")[:80]
        filename_id = filename_id or "account"
        if filename_id != sender_account_id:
            filename_id += "_" + hashlib.sha256(sender_account_id.encode("utf-8")).hexdigest()[:16]
        path = self.root / "execution_data" / ("执行数据_" + filename_id + ".csv")
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(("执行账号", "项目ID", "点赞次数", "收藏次数", "评论次数", "分享次数"))
        for task_id, project in state["project_accounts"].items():
            entry = project["accounts"].get(sender_account_id)
            result = entry.get("result") if entry else None
            if result is None:
                continue
            passed = {item["action"] for item in result["executed_actions"]
                      if item["status"] == "PASSED"}
            writer.writerow((sender_account_id, task_id,
                             int("Like" in passed), int("Favorite" in passed),
                             int("Comment" in passed), int("Share" in passed)))
        try:
            _atomic_bytes(path, output.getvalue().encode("utf-8-sig"))
        except OSError as exc:
            raise EngineError("EXECUTION_DATA_WRITE_FAILED", "执行数据输出失败：" + str(path)) from exc
        return path

    def _finish(self, state: Dict[str, Any], device_id: str, device: Dict[str, Any],
                status: str, actions: List[Dict[str, str]], error_code: str,
                message: str, blocked: bool = False) -> Dict[str, Any]:
        active = device["active_task"]
        saved = active.get("action_results", [])
        if len(actions) < len(saved) or actions[:len(saved)] != saved:
            raise EngineError("ACTION_RESULT_CONFLICT", "结束结果不能覆盖已经保存的动作进度")
        now = _utcnow()
        result = {"task_id": active["task_id"], "run_token": active["run_token"],
                  "device_id": device_id, "sender_account_id": device_id,
                  "physical_device_id": active.get("physical_device_id"),
                  "assignment_id": active.get("assignment_id"),
                  "status": status, "executed_actions": actions,
                  "error_code": error_code, "log_message": message,
                  "finished_at": now.isoformat(), "elapsed_seconds":
                  round(max(0.0, (now - _timestamp(active["claimed_at"])).total_seconds()), 3)}
        if blocked:
            device.update(state="BLOCKED", active_task=None, cooldown_until=None)
            seconds = 0
        else:
            seconds = self.devices.cool(device, now)
        result.update(cooldown_seconds=seconds, cooldown_until=device["cooldown_until"],
                      device_state=device["state"])
        result_message = _json(result)
        if len(result_message) > 32767:
            raise EngineError("RESULT_TOO_LONG", "结果超过 Excel 单元格长度上限")
        device["last_result"] = result
        entry = self._register_active(state, device_id, active)
        entry.update(status=status, result=result,
                     comment_sent=any(value["action"] == "Comment" and value["status"] == "PASSED"
                                      for value in actions))
        self._commit(state, self._project_update(state, active["task_id"]))
        return _copy(result)

    def get_next_task(self, physical_device_id: Optional[str] = None,
                      allocation_only: Optional[bool] = None) -> Optional[Dict[str, Any]]:
        if allocation_only is None:
            allocation_only = physical_device_id is not None
        if not isinstance(allocation_only, bool):
            raise EngineError("INVALID_PAYLOAD", "allocation_only必须为布尔值")
        with self._session() as state:
            now = _utcnow()
            device = self.devices.get(state, self.device_id, now)
            if device["state"] != "READY":
                self._save_state(state)
                return None
            try:
                task = self.scheduler.select(state, now, self.device_id, allocation_only)
            except EngineError:
                self._save_state(state)
                raise
            if task is None:
                self._save_state(state)
                return None
            project = state["project_accounts"].setdefault(task["task_id"], {
                "source_key": _source_key(task), "accounts": {}})
            assignment = project.get("assignment")
            entry = project["accounts"].get(self.device_id)
            if assignment:
                if not physical_device_id:
                    raise EngineError("PHYSICAL_DEVICE_REQUIRED", "固定分配任务需传入影刀本线程实际连接的设备标识")
                if physical_device_id != entry["physical_device_id"]:
                    raise EngineError("PHYSICAL_DEVICE_MISMATCH", "本线程设备与项目分配设备不一致")
            if physical_device_id:
                if not isinstance(physical_device_id, str) or not physical_device_id.strip():
                    raise EngineError("INVALID_DEVICE_ID", "实际设备标识必须为非空字符串")
                for account, other in state["devices"].items():
                    active = other.get("active_task")
                    if account != self.device_id and active and active.get("physical_device_id") == physical_device_id:
                        raise EngineError("PHYSICAL_DEVICE_BUSY", "同一物理设备正在被另一个执行账号使用")
                task["physical_device_id"] = physical_device_id
            reserved = {entry["comment_content"] for entry in project["accounts"].values()
                        if entry.get("comment_content")}
            available = [comment for comment in task["comments"] if comment not in reserved]
            if not assignment and not available:
                raise EngineError("COMMENTS_EXHAUSTED", "项目 " + task["task_id"] +
                                  " 没有未分配评论，请先补充Comments")
            task["comment_content"] = entry["comment_content"] if assignment else secrets.choice(available)
            active = self.devices.claim(device, task, now)
            self._set_workspace(active)
            self._register_active(state, self.device_id, active)
            self._commit(state, self._project_update(state, task["task_id"]))
            return _copy(active)

    def get_active_task(self, physical_device_id: Optional[str] = None) -> Optional[Dict[str, Any]]:
        """Read surviving context without replaying uncertain mobile actions."""
        with self._session() as state:
            device = self.devices.get(state, self.device_id, _utcnow())
            active = device.get("active_task")
            if active is not None and device["state"] == "BUSY":
                if active.get("assignment_id") and physical_device_id != active.get("physical_device_id"):
                    raise EngineError("PHYSICAL_DEVICE_MISMATCH", "恢复任务必须核对本线程实际连接的物理设备")
                self._set_workspace(active)
                workbook, sheet, columns, rows, _ = self.scheduler.read()
                try:
                    row = rows.get(active["task_id"])
                    status = (_text(sheet.cell(row, columns["Status"]).value).upper()
                              if row is not None else None)
                    if row is None or status != "RUNNING":
                        project = state["project_accounts"].get(active["task_id"])
                        if active.get("strategy") is not None or (project and len(project["accounts"]) > 1):
                            raise EngineError("TASK_SOURCE_CHANGED", "活动任务已删除或状态被改，且已有策略或多账号进度；请先核对")
                        # Preserve the existing pre-interaction manual reset behavior.
                        state["project_accounts"].pop(active["task_id"], None)
                        device.update(state="READY", active_task=None, cooldown_until=None)
                        self._save_state(state)
                        return None
                    refreshed = dict(active,
                        keyword=_text(sheet.cell(row, columns["Query_Keyword"]).value),
                        expected_author=_text(sheet.cell(row, columns["Expected_Author"]).value),
                        expected_douyin_id=(_douyin_id(sheet.cell(row, columns["Expected_Douyin_ID"]).value)
                                           if "Expected_Douyin_ID" in columns else ""),
                        content_type=_text(sheet.cell(row, columns["Content_Type"]).value),
                        comments=_comment_pool(sheet.cell(row, columns["Comments"]).value))
                    refreshed["search_section"] = refreshed["content_type"]
                    _validate_route(refreshed["content_type"], refreshed["content_type"], columns)
                    if not refreshed["keyword"] or not refreshed["expected_author"]:
                        raise EngineError("INVALID_TASK", "关键词和预期作者必须非空")
                    changed = (_source_key(refreshed) != _source_key(active)
                               or active.get("comment_content") not in refreshed["comments"])
                    if active.get("assignment_id"):
                        count = (_target_count(sheet.cell(row, columns["Target_Device_Count"]).value)
                                 if "Target_Device_Count" in columns else None)
                        if changed or count != active["target_device_count"]:
                            raise EngineError("ASSIGNMENT_CHANGED", "固定分配的作品、设备数量或预留评论被更改")
                    if changed and active.get("strategy") is not None:
                        raise EngineError("TASK_SOURCE_CHANGED", "当前作品或已分配评论被更改，且已有互动策略；请先核对进度")
                    self._register_active(state, self.device_id, refreshed)
                    device["active_task"] = refreshed
                    if changed:
                        project = state["project_accounts"][active["task_id"]]
                        other = [entry for account, entry in project["accounts"].items()
                                 if account != self.device_id]
                        if other and project["source_key"] != _source_key(refreshed):
                            raise EngineError("TASK_SOURCE_CHANGED", "多账号项目的作品信息不能在执行中更改")
                        project["source_key"] = _source_key(refreshed)
                finally:
                    workbook.close()
            self._save_state(state)
            return _copy(device["active_task"])

    def get_device_status(self) -> Dict[str, Any]:
        with self._session() as state:
            now = _utcnow()
            device = self.devices.get(state, self.device_id, now)
            remaining = (max(0.0, (_timestamp(device["cooldown_until"]) - now).total_seconds())
                         if device["state"] == "COOLDOWN" else 0.0)
            self._save_state(state)
            return dict(_copy(device), device_id=self.device_id,
                        cooldown_remaining_seconds=round(remaining, 3))

    def assert_author(self, expected: Any, actual: Any) -> str:
        try:
            return self.authors.assert_author(expected, actual)
        except EngineError as exc:
            self.logger.error("author_assertion_error code=%s message=%s", exc.code, exc)
            raise

    def generate_strategy(self, seed: Optional[Union[int, str]] = None,
                          task_id: Optional[str] = None,
                          run_token: Optional[str] = None) -> Dict[str, Any]:
        with self._session() as state:
            device = self.devices.get(state, self.device_id, _utcnow())
            if device["state"] in ("COOLDOWN", "BLOCKED"):
                raise EngineError("DEVICE_UNAVAILABLE", "设备冷却或熔断期间不生成执行计划")
            active = device["active_task"]
            if task_id is not None or run_token is not None:
                if task_id is None:
                    raise EngineError("INVALID_TASK_ID", "传入令牌时必须同时提供 task_id")
                active = self._owned(device, _task_id(task_id), run_token)
            if active is None:
                raise EngineError("TASK_REQUIRED", "先领取当前账号的任务，再生成互动计划")
            if active["strategy"] is not None:
                plan = active["strategy"]
                if plan.get("policy_version") != 2:
                    raise EngineError("LEGACY_STRATEGY", "已有旧版互动计划，请先核对已执行动作，不能自动改为新计划")
                if seed is not None and str(seed) != plan["seed"]:
                    raise EngineError("STRATEGY_CONFLICT", "当前任务已经保存了不同种子的计划")
                return _copy(plan)
            plan = self.strategy.generate(seed, active["comment_content"])
            plan.update(task_id=active["task_id"], run_token=active["run_token"],
                        sender_account_id=self.device_id, content_type=active["content_type"])
            active["strategy"] = plan
            entry = state["project_accounts"].get(active["task_id"], {}).get("accounts", {}).get(self.device_id)
            if entry is not None:
                entry["plan"] = _copy(plan)
            state["exports_dirty"] = True
            self._save_state(state)
            self._flush_exports(state)
            return _copy(plan)

    def record_result(self, task_id: Union[str, int], status: str,
                      executed_actions: List[Any], error_code: str = "",
                      log_message: str = "", run_token: Optional[str] = None) -> Dict[str, Any]:
        identity = _task_id(task_id)
        status = _text(status).upper()
        if status not in ("SUCCESS", "NOT_FOUND", "FAILED"):
            raise EngineError("INVALID_RESULT_STATUS", "只允许 SUCCESS / NOT_FOUND / FAILED")
        actions = self._actions(executed_actions)
        if not isinstance(error_code, str) or not isinstance(log_message, str):
            raise EngineError("INVALID_RESULT", "error_code 与 log_message 必须为字符串")
        if len(error_code) > 256 or len(log_message) > 8000:
            raise EngineError("INVALID_RESULT", "error_code 超过 256 字符或 log_message 超过 8000 字符")
        if status != "FAILED" and error_code:
            raise EngineError("INVALID_RESULT", "通过或未找到结果不能携带异常错误码")
        if status == "NOT_FOUND" and actions:
            raise EngineError("INVALID_RESULT", "未命中作者时不得有已执行交互")
        error_code = error_code or ("TEST_FAILED" if status == "FAILED" else "")
        with self._session() as state:
            device = self.devices.get(state, self.device_id, _utcnow())
            entry = state["project_accounts"].get(identity, {}).get("accounts", {}).get(self.device_id)
            last = entry.get("result") if entry else device.get("last_result")
            if last and last["task_id"] == identity and last["run_token"] == run_token:
                expected = (status, actions, error_code, log_message)
                stored = (last["status"], last["executed_actions"], last["error_code"], last["log_message"])
                if expected != stored:
                    raise EngineError("RESULT_CONFLICT", "同一次运行已经提交了不同结果")
                self._write_execution_data(state, self.device_id)
                return _copy(last)
            active = self._owned(device, identity, run_token)
            saved = active.get("action_results", [])
            if len(actions) < len(saved) or actions[:len(saved)] != saved:
                raise EngineError("ACTION_RESULT_CONFLICT", "最终结果不能覆盖已保存的动作进度")
            if status == "SUCCESS":
                plan = active["strategy"]
                if not plan:
                    raise EngineError("STRATEGY_REQUIRED", "SUCCESS 必须对应已保存的交互计划")
                if plan.get("policy_version") != 2 or not set(REQUIRED_ACTIONS).issubset(plan["execution_order"]):
                    raise EngineError("LEGACY_STRATEGY", "SUCCESS必须对应固定点赞、评论、收藏的新计划")
                if ([item["action"] for item in actions] != plan["execution_order"] or
                    any(item["status"] not in ("PASSED", "SKIPPED_ALREADY_ACTIVE") for item in actions)):
                    raise EngineError("ASSERTION_INCOMPLETE", "实际动作顺序或结果与计划不一致")
            blocked = error_code == "SECURITY_CHALLENGE_DETECTED"
            if blocked:
                device["block_info"] = {"blocked_at": _utcnow().isoformat(),
                    "log_message": log_message, "executed_actions": actions}
            return self._finish(state, self.device_id, device, status, actions,
                                error_code, log_message, blocked)

    def trigger_security_block(self, log_message: str = "检测到人机验证挑战",
                               executed_actions: Optional[List[Any]] = None) -> Dict[str, Any]:
        actions = self._actions([] if executed_actions is None else executed_actions)
        if not isinstance(log_message, str) or len(log_message) > 8000:
            raise EngineError("INVALID_RESULT", "熔断说明须为不超过 8000 字符的字符串")
        with self._session(recover=False) as state:
            device = self.devices.get(state, self.device_id, _utcnow())
            if device.get("active_task"):
                saved = device["active_task"].get("action_results", [])
                if executed_actions is None:
                    actions = _copy(saved)
                elif len(actions) < len(saved) or actions[:len(saved)] != saved:
                    raise EngineError("ACTION_RESULT_CONFLICT", "锁定结果不能覆盖已经保存的动作进度")
            if device["state"] != "BLOCKED":
                device.update(state="BLOCKED", cooldown_until=None, block_info={
                    "blocked_at": _utcnow().isoformat(), "log_message": log_message,
                    "executed_actions": actions})
                self._save_state(state)  # Emergency gate survives even an Excel write failure.
                self.logger.critical("security_block device_id=%s reason=%s", self.device_id, log_message)
            self._recover(state)
            return dict(_copy(device), device_id=self.device_id)

    def reset_blocked_device(self, operator: str, reason: str) -> Dict[str, Any]:
        """Call only after an explicit human decision; never schedule this method."""
        if not isinstance(operator, str) or not operator.strip() or not isinstance(reason, str) or not reason.strip():
            raise EngineError("MANUAL_RESET_REQUIRED", "人工复位必须提供操作人和原因")
        with self._session() as state:
            device = self.devices.get(state, self.device_id, _utcnow())
            if device["state"] != "BLOCKED":
                raise EngineError("DEVICE_NOT_BLOCKED", "当前设备不处于 BLOCKED")
            seconds = self.devices.cool(device, _utcnow())
            device["block_info"] = None
            device["last_manual_reset"] = {"operator": operator, "reason": reason,
                                           "reset_at": _utcnow().isoformat()}
            self._save_state(state)
            self.logger.warning("manual_reset device_id=%s operator=%s reason=%s", self.device_id, operator, reason)
            return {"device_id": self.device_id, "state": "COOLDOWN", "cooldown_seconds": seconds,
                    "cooldown_until": device["cooldown_until"]}


def initialize_project(root: str) -> Dict[str, Any]:
    """Create missing inputs only; never overwrite existing assets or state."""
    gateway = UnifiedGateway(root)
    created: List[str] = []
    with gateway._session(recover=False):
        files = {gateway.root / "config" / "settings.json": _json(DEFAULTS) + "\n",
                 gateway.root / "data" / "author_whitelist.json": '{"authors":[]}\n'}
        for path, content in files.items():
            if not path.exists():
                _atomic_bytes(path, content.encode("utf-8"))
                created.append(str(path))
        if not gateway.scheduler.path.exists():
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "TaskCases"
            sheet.append(list(HEADERS + ("Expected_Douyin_ID", "Content_Type", "Comments", "Target_Device_Count")))
            sheet.freeze_panes = "A2"
            for column, width in zip("ABCDEFGHIJ", (20, 40, 30, 25, 18, 40, 70, 28, 22, 48)):
                sheet.column_dimensions[column].width = width
            sheet.column_dimensions["H"].number_format = "@"
            sheet.column_dimensions["I"].number_format = "@"
            sheet.column_dimensions["J"].number_format = "@"
            sheet.column_dimensions["K"].width = 24
            output = io.BytesIO()
            workbook.save(output)
            workbook.close()
            _atomic_bytes(gateway.scheduler.path, output.getvalue())
            created.append(str(gateway.scheduler.path))
    return {"root": str(gateway.root), "created": created}


def get_next_task(root: Optional[str] = None, device_id: str = "default-device",
                  physical_device_id: Optional[str] = None) -> Optional[Dict[str, Any]]:
    return UnifiedGateway(root, device_id).get_next_task(physical_device_id)


def assert_author(expected: Any, actual: Any, root: Optional[str] = None) -> str:
    return UnifiedGateway(root).assert_author(expected, actual)


def generate_strategy(root: Optional[str] = None, device_id: str = "default-device",
                      seed: Optional[Union[int, str]] = None, task_id: Optional[str] = None,
                      run_token: Optional[str] = None) -> Dict[str, Any]:
    return UnifiedGateway(root, device_id).generate_strategy(seed, task_id, run_token)


def record_result(task_id: Union[str, int], status: str, executed_actions: List[Any],
                  error_code: str = "", log_message: str = "", run_token: Optional[str] = None,
                  root: Optional[str] = None, device_id: str = "default-device") -> Dict[str, Any]:
    return UnifiedGateway(root, device_id).record_result(task_id, status, executed_actions,
                                                       error_code, log_message, run_token)


def trigger_security_block(root: Optional[str] = None, device_id: str = "default-device",
                           log_message: str = "检测到人机验证挑战",
                           executed_actions: Optional[List[Any]] = None) -> Dict[str, Any]:
    return UnifiedGateway(root, device_id).trigger_security_block(log_message, executed_actions)


def main() -> int:
    parser = argparse.ArgumentParser(description="PhoneWorkbench 1.0 核心策略与数据引擎")
    parser.add_argument("--root", default=None)
    parser.add_argument("--device-id", default="default-device")
    parser.add_argument("--init", action="store_true")
    parser.add_argument("--operation", default="get_device_status", choices=(
        "get_next_task", "get_active_task", "get_device_status", "assert_author",
        "generate_strategy", "record_result", "trigger_security_block", "reset_blocked_device",
        "allocate_task", "get_project_progress", "save_action_results", "close_pending_assignment"))
    parser.add_argument("--payload", default="{}", help="JSON 对象，作为方法关键字参数")
    args = parser.parse_args()
    try:
        gateway = UnifiedGateway(args.root, args.device_id)
        payload = json.loads(args.payload)
        if not isinstance(payload, dict):
            raise EngineError("INVALID_PAYLOAD", "payload 必须为 JSON 对象")
        data = (initialize_project(str(gateway.root)) if args.init else
                getattr(gateway, args.operation)(**payload))
        print(_json({"ok": True, "data": data}))
        return 0
    except Exception as exc:
        print(_json({"ok": False, "error_code": getattr(exc, "code", "ENGINE_ERROR"),
                     "message": str(exc)}))
        return 1


if __name__ == "__main__":
    sys.exit(main())


