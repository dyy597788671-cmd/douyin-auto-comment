import json
import unittest
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
import core_runner as core
import module1 as bridge
from test_execution_output import temporary_root, prepare_root, finish_kwargs, read_rows


def setup(root, count=2, tasks=1):
    prepare_root(root, [])
    path = root / 'data' / 'task_cases.xlsx'
    wb = load_workbook(path)
    sheet = wb.active
    sheet.cell(1, 11, 'Target_Device_Count')
    sheet.cell(2, 10, 'a\nb\nc\nd\ne')
    sheet.cell(2, 11, count)
    for index in range(1, tasks):
        sheet.append([str(1005 + index), '关键词', '作者', 0, 'PENDING', None, None,
                      '12345', '视频', 'a\nb\nc\nd\ne', count])
    wb.save(path)
    wb.close()
    return [{'udid': 'usb-' + str(index), 'custom_name': 'label-' + str(index)}
            for index in range(count)]


def read(binding, wait=True):
    return bridge.read_task(sender_account_id=binding['sender_account_id'],
                            physical_device_id=binding['device_id'], wait_for_task=wait)


def complete(root, task, status='SUCCESS'):
    gateway = core.UnifiedGateway(str(root), task['sender_account_id'])
    actions = []
    if status == 'SUCCESS':
        plan = gateway.generate_strategy(seed=17, task_id=task['task_id'], run_token=task['run_token'])
        actions = [{'action': name, 'status': 'PASSED', 'message': ''} for name in plan['execution_order']]
        gateway.save_action_results(task['task_id'], task['run_token'], actions)
    return gateway.record_result(**finish_kwargs(task, actions, status))


class DeviceEntryTests(unittest.TestCase):
    def test_switch_device_same_task_has_independent_results(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            devices = bridge.prepare_devices(setup(root))
            first = read(devices[0])
            complete(root, first)
            second = read(devices[1])
            self.assertEqual(first['task_id'], second['task_id'])
            self.assertNotEqual(first['sender_account_id'], second['sender_account_id'])
            self.assertNotEqual(first['comment_content'], second['comment_content'])
            self.assertNotEqual(first['workspace_path'], second['workspace_path'])
            complete(root, second)
            for device in devices:
                self.assertEqual(len(read_rows(root, device['sender_account_id'])), 1)
            self.assertEqual(bridge.prepare_devices(setup_infos(2)), [])

    def test_new_device_and_reordered_connections_need_no_registration(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            infos = setup(root, 2)
            old = bridge.prepare_devices(infos)
            path = root / 'data' / 'task_cases.xlsx'
            wb = load_workbook(path)
            wb.active.append(['2000', '新作品', '作者', 0, 'PENDING', None, None,
                              '12345', '图文', 'a\nb\nc', 3])
            wb.save(path)
            wb.close()
            infos = [{'udid': 'new-phone'}] + list(reversed(infos))
            devices = bridge.prepare_devices(infos)
            self.assertEqual(len(devices), 3)
            for device in devices:
                self.assertEqual(infos[device['connection_index']]['udid'], device['device_id'])
            self.assertEqual({v['sender_account_id'] for v in old},
                             {v['sender_account_id'] for v in devices if v['device_id'] != 'new-phone'})

    def test_random_assignment_uses_only_selected_devices(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            infos = setup(root, 3)
            infos += [{'udid': 'extra-1'}, {'udid': 'extra-2'}]
            devices = bridge.prepare_devices(infos)
            self.assertEqual(len(devices), 3)
            selected = {v['device_id'] for v in devices}
            other = next(v for v in infos if v['udid'] not in selected)
            self.assertIsNone(bridge.read_task('设备_' + other['udid'], physical_device_id=other['udid'],
                                              wait_for_task=True))

    def test_missing_count_is_an_explicit_input_error(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            infos = setup(root, 1)
            path = root / 'data' / 'task_cases.xlsx'
            wb = load_workbook(path)
            wb.active.cell(2, 11).value = None
            wb.save(path)
            wb.close()
            with self.assertRaises(core.EngineError) as raised:
                bridge.prepare_devices(infos)
            self.assertEqual(raised.exception.code, 'TARGET_COUNT_REQUIRED')

    def test_empty_input_has_no_workers(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            infos = setup(root, 1)
            path = root / 'data' / 'task_cases.xlsx'
            wb = load_workbook(path)
            wb.active.delete_rows(2)
            wb.save(path)
            wb.close()
            self.assertEqual(bridge.prepare_devices(infos), [])

    def test_cooldown_and_schedule_wait_without_ending_worker(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            infos = setup(root, 1, 2)
            config = root / 'config' / 'settings.json'
            settings = json.loads(config.read_text())
            settings.update(cooldown_min_seconds=65, cooldown_max_seconds=65)
            config.write_text(json.dumps(settings))
            clock = [core._utcnow()]
            waits = []
            def advance(seconds):
                # Another account can acquire the project lock during the wait.
                core.UnifiedGateway(folder, 'other-device').get_device_status()
                waits.append(seconds)
                clock[0] += timedelta(seconds=seconds)
            with patch.object(core, '_utcnow', side_effect=lambda: clock[0]), patch('time.sleep', advance):
                device = bridge.prepare_devices(infos)[0]
                first = read(device)
                complete(root, first)
                path = root / 'data' / 'task_cases.xlsx'
                wb = load_workbook(path)
                wb.active.cell(3, 4, (clock[0] + timedelta(seconds=95)).isoformat())
                wb.save(path)
                wb.close()
                second = read(device)
                self.assertEqual(second['task_id'], '1006')
                self.assertAlmostEqual(sum(waits), 95)
                complete(root, second)
                self.assertIsNone(read(device))
                self.assertEqual(sum(waits), 95)

    def test_restart_keeps_assignment_and_only_remaining_device(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            infos = setup(root)
            devices = bridge.prepare_devices(infos)
            complete(root, read(devices[0]))
            remaining = bridge.prepare_devices(list(reversed(infos)))
            self.assertEqual(len(remaining), 1)
            self.assertEqual(remaining[0]['device_id'], devices[1]['device_id'])

    def test_claim_exception_is_reported_as_active_not_empty_on_restart(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            infos = setup(root, 1)
            device = bridge.prepare_devices(infos)[0]
            original = core.UnifiedGateway._flush_exports
            injected = [False]
            def fail_after_claim(gateway, state):
                if (not injected[0] and state['exports_dirty'] and
                        state['devices'].get(gateway.device_id, {}).get('active_task')):
                    injected[0] = True
                    raise OSError('simulated output error')
                return original(gateway, state)
            with patch.object(core.UnifiedGateway, '_flush_exports', fail_after_claim):
                with self.assertRaises(core.EngineError):
                    read(device)
            restarted = bridge.prepare_devices(infos)
            self.assertEqual(len(restarted), 1)
            with self.assertRaisesRegex(RuntimeError, '已领取'):
                read(restarted[0])

    def test_not_found_is_recorded_for_this_device_only(self):
        with temporary_root() as folder, patch.object(bridge, 'PROJECT_ROOT', folder):
            root = Path(folder)
            devices = bridge.prepare_devices(setup(root))
            first = read(devices[0])
            complete(root, first, 'NOT_FOUND')
            self.assertEqual(read(devices[1])['task_id'], first['task_id'])


def setup_infos(count):
    return [{'udid': 'usb-' + str(index)} for index in range(count)]


if __name__ == '__main__':
    unittest.main()
