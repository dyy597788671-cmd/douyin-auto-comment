import json
import multiprocessing
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import core_runner as core
from openpyxl import load_workbook
from test_execution_output import temporary_root, prepare_root, finish_kwargs, read_rows


ROSTER = [{"device_id": "usb-" + str(i), "sender_account_id": "手机0" + str(i)}
          for i in range(1, 6)]
ACCOUNTS = [v["sender_account_id"] for v in ROSTER]


def setup(root, count=4):
    prepare_root(root, ACCOUNTS)
    workbook = load_workbook(root / "data" / "task_cases.xlsx")
    workbook.active.cell(1, 11, "Target_Device_Count")
    workbook.active.cell(2, 11, count)
    workbook.save(root / "data" / "task_cases.xlsx")
    workbook.close()
    return core.UnifiedGateway(str(root), ACCOUNTS[0])


def edit_cell(root, column, value, row=2):
    path = root / "data" / "task_cases.xlsx"
    workbook = load_workbook(path)
    workbook.active.cell(row, column).value = value
    workbook.save(path)
    workbook.close()


def claim(root, value):
    gateway = core.UnifiedGateway(str(root), value["sender_account_id"])
    task = gateway.get_next_task(physical_device_id=value["device_id"])
    if task is None:
        return gateway, None, []
    plan = gateway.generate_strategy(seed=17, task_id=task["task_id"], run_token=task["run_token"])
    actions = [{"action": name, "status": "PASSED", "message": ""}
               for name in plan["execution_order"]]
    return gateway, task, actions


def allocation_worker(folder, value):
    gateway = core.UnifiedGateway(folder, value["sender_account_id"])
    gateway.allocate_task("1005", ROSTER)
    gateway, task, actions = claim(Path(folder), value)
    if task:
        gateway.save_action_results(task["task_id"], task["run_token"], actions)
        gateway.record_result(**finish_kwargs(task, actions))


class DeviceAllocationTests(unittest.TestCase):
    def expect_code(self, code, call):
        with self.assertRaises(core.EngineError) as raised:
            call()
        self.assertEqual(raised.exception.code, code)

    def test_five_connections_choose_four_and_freeze_on_restart(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root)
            first = gateway.allocate_task("1005", ROSTER)
            self.assertEqual(first["assigned_device_count"], 4)
            self.assertEqual(first["finished_device_count"], 0)
            self.assertFalse(first["completed"])
            self.assertEqual(len({v["physical_device_id"] for v in first["accounts"]}), 4)
            state = json.loads(gateway.state_path.read_text())
            entries = state["project_accounts"]["1005"]["accounts"]
            self.assertEqual(len({v["comment_content"] for v in entries.values()}), 4)
            restarted = core.UnifiedGateway(folder, ACCOUNTS[0])
            second = restarted.allocate_task("1005", list(reversed(ROSTER)))
            self.assertEqual(first, second)
            # A later connection snapshot does not replace the saved participants.
            third = restarted.allocate_task("1005", ROSTER[:1])
            self.assertEqual(first, third)

    def test_only_assigned_workers_and_complete_at_four(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root)
            allocation = gateway.allocate_task("1005", ROSTER)
            selected = {v["sender_account_id"] for v in allocation["accounts"]}
            finished = 0
            for value in ROSTER:
                worker, task, actions = claim(root, value)
                if value["sender_account_id"] not in selected:
                    self.assertIsNone(task)
                    continue
                self.assertEqual(task["physical_device_id"], value["device_id"])
                worker.record_result(**finish_kwargs(task, actions))
                finished += 1
                progress = gateway.get_project_progress("1005")
                self.assertEqual(progress["finished_device_count"], finished)
                self.assertEqual(progress["completed"], finished == 4)
                self.assertEqual(progress["status"], "SUCCESS" if finished == 4 else "RUNNING")
                self.assertIsNone(worker.get_next_task(value["device_id"]))
            self.assertEqual(progress["counts"], dict.fromkeys(core.ACTIONS, 4))
            self.assertEqual(progress["target_device_count"], 4)
            exported = json.loads((root / "execution_data" / "project_progress.json").read_text())
            self.assertEqual(exported["projects"][0], progress)

    def test_partial_failure_and_not_found_finish_without_fake_counts(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 3)
            allocation = gateway.allocate_task("1005", ROSTER)
            bindings = {v["sender_account_id"]: v for v in ROSTER}
            for i, detail in enumerate(allocation["accounts"]):
                worker, task, actions = claim(root, bindings[detail["sender_account_id"]])
                if i == 0:
                    worker.record_result(**finish_kwargs(task, actions))
                elif i == 1:
                    actions[-1]["status"] = "FAILED"
                    worker.save_action_results(task["task_id"], task["run_token"], actions)
                    worker.record_result(**finish_kwargs(task, actions, "FAILED"))
                else:
                    worker.record_result(**finish_kwargs(task, [], "NOT_FOUND"))
            progress = gateway.get_project_progress("1005")
            self.assertTrue(progress["completed"])
            self.assertEqual(progress["status"], "FAILED")
            self.assertEqual(progress["counts"], {"Like": 2, "Comment": 2, "Favorite": 2, "Share": 1})
            self.assertEqual([progress[v] for v in ("success_device_count", "failed_device_count", "not_found_device_count")], [1, 1, 1])

    def test_count_greater_than_connections_rejected_before_reservation(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 6)
            self.expect_code("TARGET_COUNT_EXCEEDS_CONNECTED", lambda: gateway.allocate_task("1005", ROSTER))
            self.assertFalse(gateway.state_path.exists())

    def test_invalid_counts_and_duplicate_devices(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root)
            for count in (0, -1, True, 1.5, "四", "4.0"):
                edit_cell(root, 11, count)
                self.expect_code("INVALID_TARGET_COUNT", lambda: gateway.allocate_task("1005", ROSTER))
            edit_cell(root, 11, 4)
            self.expect_code("DUPLICATE_DEVICE_ACCOUNT", lambda: gateway.allocate_task("1005", ROSTER + [ROSTER[0]]))
            self.expect_code("CONNECTED_DEVICES_REQUIRED", lambda: gateway.allocate_task("1005", []))

    def test_no_connected_roster_or_allocation_no_claim(self):
        with temporary_root() as folder:
            gateway = setup(Path(folder))
            self.expect_code("INVALID_TASK_DATA", gateway.get_next_task)
            self.assertFalse(gateway.state_path.read_text().find('"assignment"') >= 0)

    def test_comment_capacity_uses_selected_count(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root)
            edit_cell(root, 10, "a\nb\nc")
            self.expect_code("COMMENTS_INSUFFICIENT", lambda: gateway.allocate_task("1005", ROSTER))
            edit_cell(root, 10, "a\nb\nc\nd")
            gateway.allocate_task("1005", ROSTER)
            # Four distinct candidates suffice for a five-device pool.
            for value in ROSTER:
                worker, task, actions = claim(root, value)
                if task:
                    worker.record_result(**finish_kwargs(task, actions))
            self.assertTrue(gateway.get_project_progress("1005")["completed"])

    def test_claim_and_restore_require_matching_physical_device(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            self.expect_code("PHYSICAL_DEVICE_REQUIRED", gateway.get_next_task)
            self.expect_code("PHYSICAL_DEVICE_MISMATCH", lambda: gateway.get_next_task("other-usb"))
            task = gateway.get_next_task("usb-1")
            self.expect_code("PHYSICAL_DEVICE_MISMATCH", gateway.get_active_task)
            self.expect_code("PHYSICAL_DEVICE_MISMATCH", lambda: gateway.get_active_task("other-usb"))
            self.assertEqual(task["run_token"], gateway.get_active_task("usb-1")["run_token"])

    def test_action_journal_survives_and_refuses_overwrite(self):
        with temporary_root() as folder:
            root = Path(folder)
            controller = setup(root, 1)
            controller.allocate_task("1005", ROSTER[:1])
            worker, task, actions = claim(root, ROSTER[0])
            worker.save_action_results(task["task_id"], task["run_token"], actions[:1])
            restarted = core.UnifiedGateway(folder, ACCOUNTS[0])
            active = restarted.get_active_task("usb-1")
            self.assertEqual(active["action_results"], actions[:1])
            self.assertEqual(sum(controller.get_project_progress("1005")["counts"].values()), 1)
            self.expect_code("ACTION_RESULT_CONFLICT", lambda: restarted.save_action_results("1005", task["run_token"], []))
            changed = [dict(actions[0], status="FAILED")]
            self.expect_code("ACTION_RESULT_CONFLICT", lambda: restarted.save_action_results("1005", task["run_token"], changed))
            self.expect_code("ACTION_RESULT_CONFLICT", lambda: restarted.record_result(**finish_kwargs(task, [], "FAILED")))
            self.assertFalse(controller.get_project_progress("1005")["completed"])
            worker.record_result(**finish_kwargs(task, actions))

    def test_changed_count_source_and_reserved_comment_rejected(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root)
            gateway.allocate_task("1005", ROSTER)
            for column, value, restore in ((11, 3, 4), (2, "另一个作品", "关键词"), (10, "只有一条", "\n".join("评论" + str(i) for i in range(5)))):
                edit_cell(root, column, value)
                self.expect_code("ASSIGNMENT_CHANGED", lambda: gateway.allocate_task("1005", ROSTER))
                edit_cell(root, column, restore)

    def test_pending_operator_closure_no_replacement_and_zero_counts(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            original = gateway.allocate_task("1005", ROSTER[:1])
            progress = gateway.close_pending_assignment("1005", ACCOUNTS[0], "操作人", "该设备无法继续")
            self.assertTrue(progress["completed"])
            self.assertEqual(progress["counts"], dict.fromkeys(core.ACTIONS, 0))
            self.assertEqual(progress["assignment_id"], original["assignment_id"])
            self.assertIsNone(gateway.get_next_task("usb-1"))
            self.expect_code("ASSIGNMENT_NOT_PENDING", lambda: gateway.close_pending_assignment("1005", ACCOUNTS[0], "操作人", "再次结束"))

    def test_blocked_connected_device_not_selected(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 5)
            blocked = core.UnifiedGateway(folder, ACCOUNTS[0])
            blocked.trigger_security_block()
            self.expect_code("ELIGIBLE_DEVICES_INSUFFICIENT", lambda: gateway.allocate_task("1005", ROSTER))
            edit_cell(root, 11, 4)
            allocation = gateway.allocate_task("1005", ROSTER)
            self.assertNotIn(ACCOUNTS[0], [v["sender_account_id"] for v in allocation["accounts"]])

    def test_excel_write_failure_recovers_same_assignment(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root)
            with patch.object(core.TaskScheduler, "apply_update", side_effect=PermissionError("occupied")):
                self.expect_code("WRITE_PENDING", lambda: gateway.allocate_task("1005", ROSTER))
            state = json.loads(gateway.state_path.read_text())
            assignment_id = state["project_accounts"]["1005"]["assignment"]["assignment_id"]
            progress = gateway.allocate_task("1005", list(reversed(ROSTER)))
            self.assertEqual(progress["assignment_id"], assignment_id)
            self.assertIsNone(json.loads(gateway.state_path.read_text())["pending_transaction"])

    def test_output_failure_repairs_without_repeat_actions(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            worker, task, actions = claim(root, ROSTER[0])
            write = core._atomic_bytes
            def fail_output(path, payload):
                if path.name == "project_progress.json":
                    raise PermissionError("occupied")
                write(path, payload)
            with patch.object(core, "_atomic_bytes", side_effect=fail_output):
                self.expect_code("EXECUTION_DATA_WRITE_FAILED", lambda: worker.record_result(**finish_kwargs(task, actions)))
            self.assertTrue(json.loads(gateway.state_path.read_text())["exports_dirty"])
            self.assertTrue(gateway.get_project_progress("1005")["completed"])
            self.assertFalse(json.loads(gateway.state_path.read_text())["exports_dirty"])
            worker.record_result(**finish_kwargs(task, actions))
            self.assertEqual(len(read_rows(root, ACCOUNTS[0])), 1)

    def test_five_processes_allocate_once_and_only_four_finish(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root)
            context = multiprocessing.get_context("spawn")
            processes = [context.Process(target=allocation_worker, args=(folder, value)) for value in ROSTER]
            for process in processes:
                process.start()
            for process in processes:
                process.join(30)
                if process.is_alive():
                    process.terminate()
                    process.join()
                self.assertEqual(process.exitcode, 0)
            progress = gateway.get_project_progress("1005")
            self.assertTrue(progress["completed"])
            self.assertEqual(progress["finished_device_count"], 4)
            self.assertEqual(progress["counts"], dict.fromkeys(core.ACTIONS, 4))
            self.assertEqual(len(list((root / "execution_data").glob("*.csv"))), 4)

    def test_late_result_retry_after_another_project(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            worker, first_task, first_actions = claim(root, ROSTER[0])
            first_result = worker.record_result(**finish_kwargs(first_task, first_actions))
            path = root / "data" / "task_cases.xlsx"
            workbook = load_workbook(path)
            workbook.active.append(["1006", "关键词", "作者", 0, "PENDING", None, None,
                                    "12345", "视频", "新评论", 1])
            workbook.save(path)
            workbook.close()
            gateway.allocate_task("1006", ROSTER[:1])
            worker, second_task, second_actions = claim(root, ROSTER[0])
            worker.record_result(**finish_kwargs(second_task, second_actions))
            self.assertEqual(worker.record_result(**finish_kwargs(first_task, first_actions)), first_result)
            self.assertEqual(len(read_rows(root, ACCOUNTS[0])), 2)
            changed = [dict(first_actions[0], status="FAILED")] + first_actions[1:]
            self.expect_code("RESULT_CONFLICT", lambda: worker.record_result(**finish_kwargs(first_task, changed, "FAILED")))

    def test_physical_device_cannot_run_two_accounts_at_once(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            worker, task, actions = claim(root, ROSTER[0])
            path = root / "data" / "task_cases.xlsx"
            workbook = load_workbook(path)
            workbook.active.append(["1006", "关键词", "作者", 0, "PENDING", None, None,
                                    "12345", "视频", "新评论", 1])
            workbook.save(path)
            workbook.close()
            new_roster = [{"device_id": "usb-1", "sender_account_id": ACCOUNTS[1]}]
            gateway.allocate_task("1006", new_roster)
            other = core.UnifiedGateway(folder, ACCOUNTS[1])
            self.expect_code("PHYSICAL_DEVICE_BUSY", lambda: other.get_next_task("usb-1"))
            self.expect_code("ASSIGNMENT_NOT_PENDING", lambda: gateway.close_pending_assignment("1005", ACCOUNTS[0], "操作人", "结束"))
            worker.record_result(**finish_kwargs(task, actions))
            self.assertEqual(other.get_next_task("usb-1")["task_id"], "1006")

    def test_security_block_preserves_confirmed_partial_counts(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            worker, task, actions = claim(root, ROSTER[0])
            worker.save_action_results("1005", task["run_token"], actions[:2])
            worker.trigger_security_block()
            progress = gateway.get_project_progress("1005")
            self.assertTrue(progress["completed"])
            self.assertEqual(progress["failed_device_count"], 1)
            self.assertEqual(sum(progress["counts"].values()), 2)
            self.assertEqual(progress["accounts"][0]["error_code"], "SECURITY_CHALLENGE_DETECTED")

    def test_wrong_token_and_corrupt_allocation_never_rebuilt(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            worker, task, actions = claim(root, ROSTER[0])
            self.expect_code("TASK_NOT_OWNED", lambda: worker.save_action_results("1005", "wrong-token", actions))
            state = json.loads(gateway.state_path.read_text())
            state["project_accounts"]["1005"]["assignment"]["selected_accounts"] = [ACCOUNTS[1]]
            gateway.state_path.write_text(json.dumps(state), encoding="utf-8")
            before = gateway.state_path.read_bytes()
            self.expect_code("STATE_CORRUPT", lambda: gateway.get_project_progress("1005"))
            self.assertEqual(gateway.state_path.read_bytes(), before)

    def test_workspace_files_are_separate_and_stable(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 5)
            gateway.allocate_task("1005", ROSTER)
            paths = set()
            for value in ROSTER:
                worker = core.UnifiedGateway(folder, value["sender_account_id"])
                task = worker.get_next_task(value["device_id"])
                paths.add(task["workspace_path"])
                image_path = Path(task["workspace_path"]) / "dy_filter_screen.png"
                image_path.write_text(value["sender_account_id"])
                restored = worker.get_active_task(value["device_id"])
                self.assertEqual(restored["workspace_path"], task["workspace_path"])
                self.assertEqual(image_path.read_text(), value["sender_account_id"])
            self.assertEqual(len(paths), 5)

    def test_bridge_refuses_duplicate_worker_before_strategy(self):
        import module1
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            with patch.object(module1, "PROJECT_ROOT", folder):
                first = module1.read_task(ACCOUNTS[0], physical_device_id="usb-1")
                with self.assertRaisesRegex(RuntimeError, "不能重复启动"):
                    module1.read_task(ACCOUNTS[0], physical_device_id="usb-1")
                restored = module1.read_task(ACCOUNTS[0], physical_device_id="usb-1", resume_active=True)
                self.assertEqual(restored["run_token"], first["run_token"])

    def test_new_workers_never_claim_unassigned_legacy_rows(self):
        with temporary_root() as folder:
            root = Path(folder)
            gateway = setup(root, 1)
            gateway.allocate_task("1005", ROSTER[:1])
            path = root / "data" / "task_cases.xlsx"
            workbook = load_workbook(path)
            workbook.active.append(["old-project", "旧关键词", "作者", 0, "PENDING", None, None,
                                    "12345", "视频", "旧评论1\n旧评论2", None])
            workbook.save(path)
            workbook.close()
            other = core.UnifiedGateway(folder, ACCOUNTS[1])
            self.assertIsNone(other.get_next_task("usb-2"))
            task = gateway.get_next_task("usb-1")
            self.assertEqual(task["task_id"], "1005")


if __name__ == "__main__":
    unittest.main()
